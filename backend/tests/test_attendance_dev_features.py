import unittest
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import numpy as np
from fastapi import HTTPException
from openpyxl import load_workbook
from PIL import Image

from app.db.database import (
    AdminUser,
    Coordinator,
    RosterEntry,
    _validated_hierarchy,
    migrate_legacy_reporting_hierarchy,
)
from app.services.attendance_export_xlsx import (
    build_coordinator_details_workbook,
    build_export_workbook,
    build_participant_details_workbook,
)
from app.services.geofence import check_geofence
from app.services.liveness import check_passive_liveness
from app.services.ocr import extract_overlay_data
from app.config import settings
from app.main import (
    AttendanceOverrideRequest,
    ManualAttendanceRequest,
    create_manual_entry,
    export_coordinator_details,
    export_employee_details,
    get_attendance,
    patch_attendance_override,
)


class OptionalAttendanceOverrideReasonTests(unittest.IsolatedAsyncioTestCase):
    async def test_status_override_accepts_missing_reason_and_keeps_actor_audit(self):
        admin = AdminUser(username="manual", password_hash="hash", role="manual_team")
        request = AttendanceOverrideRequest(
            employee_id="EMP001",
            date="2026-10-22",
            attendance_code="P",
        )
        with patch(
            "app.main.save_attendance_override",
            new_callable=AsyncMock,
            return_value={"attendance_code": "P"},
        ) as save_override:
            await patch_attendance_override(request=request, admin=admin)

        save_override.assert_awaited_once_with(
            "EMP001",
            "2026-10-22",
            "P",
            "",
            "manual",
        )


class ManualEntryStatusOverrideConflictTests(unittest.IsolatedAsyncioTestCase):
    async def test_saved_override_conflict_suggests_changing_status(self):
        class Database:
            submissions = SimpleNamespace(find_one=AsyncMock(return_value=None))
            attendance_overrides = SimpleNamespace(
                find_one=AsyncMock(return_value={"attendance_code": "A"})
            )

        database = Database()
        request = ManualAttendanceRequest(
            employee_id="EMP002",
            date="2026-10-22",
            reason="Phone died",
        )
        admin = AdminUser(username="manual", password_hash="hash", role="manual_team")
        with (
            patch(
                "app.main.create_manual_attendance",
                new_callable=AsyncMock,
                side_effect=ValueError("duplicate"),
            ),
            patch("app.main.get_db", return_value=database),
        ):
            with self.assertRaises(HTTPException) as error:
                await create_manual_entry(request=request, admin=admin)

        self.assertEqual(error.exception.status_code, 409)
        self.assertEqual(error.exception.detail["code"], "status_override_exists")
        self.assertEqual(error.exception.detail["attendance_code"], "A")
        self.assertIn("deleting an attendance record does not remove", error.exception.detail["message"])
        database.submissions.find_one.assert_awaited_once()
        database.attendance_overrides.find_one.assert_awaited_once()


class AttendanceStatusOverrideDisplayTests(unittest.IsolatedAsyncioTestCase):
    async def test_attendance_list_displays_saved_status_override(self):
        class AsyncCursor:
            def __init__(self, items):
                self.items = iter(items)

            def __aiter__(self):
                return self

            async def __anext__(self):
                try:
                    return next(self.items)
                except StopIteration:
                    raise StopAsyncIteration

        override = {
            "employee_id": "EMP001",
            "attendance_date": "2026-10-22",
            "attendance_code": "A",
        }
        database = SimpleNamespace(
            attendance_overrides=SimpleNamespace(
                find=lambda query: AsyncCursor([override])
            )
        )
        record = {
            "_id": "submission-1",
            "employee_id": "EMP001",
            "attendance_date": "2026-10-22",
            "verdict": "verified",
            "review_status": "n/a",
            "attendance_status": "present",
            "shift_status": "on_time",
            "entry_type": "whatsapp",
            "issues": [],
            "received_at_utc": "2026-10-22T05:00:00Z",
            "geofence_passed": None,
        }
        admin = AdminUser(username="admin", password_hash="hash")
        with (
            patch("app.main.list_attendance", new_callable=AsyncMock, return_value=([record], 1)),
            patch("app.main.get_db", return_value=database),
        ):
            response = await get_attendance(admin=admin)

        self.assertEqual(response["records"][0]["attendance_code"], "A")


class CoordinatorDetailsWorkbookTests(unittest.TestCase):
    def test_coordinator_details_workbook_contains_coordinator_fields_and_areas_only(self):
        workbook_bytes = build_coordinator_details_workbook(
            coordinators=[{
                "pc_id": "PC001",
                "name": "Naveen",
                "phone_number": "919977799777",
                "is_active": True,
            }],
            employees=[
                {"pc_id": "PC001", "area": "Nipahd"},
                {"pc_id": "PC001", "area": "Asthi"},
                {"pc_id": "PC001", "area": "Asthi"},
            ],
        )
        workbook = load_workbook(BytesIO(workbook_bytes), data_only=True)
        sheet = workbook["Project Coordinators"]

        self.assertEqual(
            [sheet.cell(1, column).value for column in range(1, 7)],
            ["S.No", "ID", "Name", "Phone", "Status", "Participant Areas"],
        )
        self.assertEqual(
            [sheet.cell(2, column).value for column in range(1, 7)],
            [1, "PC001", "Naveen", "919977799777", "Active", "Asthi, Nipahd"],
        )
        self.assertEqual(sheet["A4"].value, "Total project co-ordinators")
        self.assertEqual(sheet["B4"].value, 1)
        self.assertEqual(workbook.sheetnames, ["Project Coordinators"])


class CoordinatorDetailsExportEndpointTests(unittest.IsolatedAsyncioTestCase):
    async def test_export_returns_filtered_coordinator_details_xlsx(self):
        coordinator = Coordinator(
            pc_id="PC001",
            name="Naveen",
            phone_number="919977799777",
        )
        employee = RosterEntry(
            phone_number="911234567890",
            employee_id="EMP001",
            name="Charan",
            role="FS",
            pc_id="PC001",
            area="Asthi",
            designation="Field Supervisor",
        )
        admin = AdminUser(username="admin", password_hash="hash")
        with (
            patch("app.main.list_coordinators", new_callable=AsyncMock, return_value=([coordinator], 1)) as list_pcs,
            patch("app.main.list_employees", new_callable=AsyncMock, return_value=([employee], 1)) as list_people,
        ):
            response = await export_coordinator_details(search="Naveen", admin=admin)

        list_pcs.assert_awaited_once_with(search="Naveen", page=1, page_size=100_000)
        list_people.assert_awaited_once_with(page=1, page_size=100_000)
        self.assertEqual(
            response.headers["content-type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(response.body), data_only=True)
        self.assertEqual(workbook.sheetnames, ["Project Coordinators"])
        self.assertEqual(workbook.active["B2"].value, "PC001")
        self.assertEqual(workbook.active["F2"].value, "Asthi")


class MultilingualOcrTests(unittest.TestCase):
    def _image_bytes(self) -> bytes:
        image = Image.new("RGB", (2, 2), "white")
        output = BytesIO()
        image.save(output, format="PNG")
        return output.getvalue()

    def test_hindi_overlay_extracts_coordinates_place_and_timestamp(self):
        text = (
            "पुणे, महाराष्ट्र, भारत\n"
            "अक्षांश: १८.५२०४° रेखांश: ७३.८५६७°\n"
            "गुरुवार, २२/१०/२०२६ १०:१५ पूर्वाह्न GMT +05:30"
        )
        with patch("app.services.ocr.pytesseract.image_to_string", return_value=text):
            result = extract_overlay_data(self._image_bytes())

        self.assertEqual(result.latitude, 18.5204)
        self.assertEqual(result.longitude, 73.8567)
        self.assertEqual(result.place, "पुणे, महाराष्ट्र, भारत")
        self.assertEqual(result.timestamp.isoformat(), "2026-10-22T10:15:00+05:30")
        self.assertTrue(result.extraction_complete)
        self.assertTrue(check_geofence(result.latitude, result.longitude, 18.5204, 73.8567)[0])

    def test_marathi_day_period_with_24_hour_time_is_not_reinterpreted(self):
        text = (
            "अक्षांश: १८.५२०४° रेखांश: ७३.८५६७°\n"
            "२२/१०/२०२६ १८:१५ दुपारी GMT +05:30"
        )
        with patch("app.services.ocr.pytesseract.image_to_string", return_value=text):
            result = extract_overlay_data(self._image_bytes())

        self.assertEqual(result.timestamp.isoformat(), "2026-10-22T18:15:00+05:30")
        self.assertTrue(result.extraction_complete)

    def test_english_overlay_remains_supported(self):
        text = (
            "Pune, Maharashtra, India\n"
            "Lat 18.5204° Long 73.8567°\n"
            "Friday, 28/08/2026 09:27 AM GMT +05:30"
        )
        with patch("app.services.ocr.pytesseract.image_to_string", return_value=text):
            result = extract_overlay_data(self._image_bytes())

        self.assertEqual(result.timestamp.isoformat(), "2026-08-28T09:27:00+05:30")
        self.assertTrue(result.extraction_complete)

    def test_gps_map_camera_ocr_punctuation_and_seconds_are_parsed(self):
        text = (
            "Lat 18. 890764., Long 73.288878\n"
            "Wednesday, 21/10/2026 07:32:14 GMT +05:30"
        )
        with patch("app.services.ocr.pytesseract.image_to_string", return_value=text):
            result = extract_overlay_data(self._image_bytes())

        self.assertEqual(result.latitude, 18.890764)
        self.assertEqual(result.longitude, 73.288878)
        self.assertEqual(result.timestamp.isoformat(), "2026-10-21T07:32:14+05:30")
        self.assertTrue(result.extraction_complete)

    def test_gps_map_camera_ocr_repairs_missing_decimal_and_misread_timezone(self):
        crop_text = (
            "Lat 18.63070+ Long 73 286878\n"
            "Wednesday, 21/10/2026 07:32 दाता +05 99"
        )
        with patch(
            "app.services.ocr.pytesseract.image_to_string",
            side_effect=["", crop_text],
        ):
            result = extract_overlay_data(self._image_bytes())

        self.assertEqual(result.latitude, 18.6307)
        self.assertEqual(result.longitude, 73.286878)
        self.assertEqual(result.timestamp.isoformat(), "2026-10-21T07:32:00+05:30")
        self.assertTrue(result.extraction_complete)


class FaceSizeReviewTests(unittest.TestCase):
    @staticmethod
    def _image_bytes() -> bytes:
        pixels = (np.indices((100, 100)).sum(axis=0) % 2 * 255).astype(np.uint8)
        image = Image.fromarray(pixels).convert("RGB")
        output = BytesIO()
        image.save(output, format="PNG")
        return output.getvalue()

    def _check_with_face_ratio(self, face_ratio: float) -> tuple[str, list[str], dict]:
        face_width = 27
        face_height = round(face_ratio * 10000 / face_width)
        face = SimpleNamespace(
            bbox=[10, 10, 10 + face_width, 10 + face_height],
            det_score=1.0,
        )
        with (
            patch.object(settings, "liveness_min_face_ratio", 0.07),
            patch("app.services.liveness.analyze_faces", return_value=[face]),
            patch("app.services.liveness.get_primary_face", return_value=face),
            patch("app.services.liveness.get_significant_faces", return_value=[face]),
        ):
            return check_passive_liveness(self._image_bytes())

    def test_seven_percent_threshold_accepts_the_observed_clear_face_size(self):
        status, issues, metrics = self._check_with_face_ratio(0.0723)

        self.assertEqual(metrics["face_ratio"], 0.0729)
        self.assertEqual(status, "passed")
        self.assertNotIn("face_too_small", issues)

    def test_faces_still_below_the_test_threshold_are_flagged(self):
        status, issues, _ = self._check_with_face_ratio(0.06)

        self.assertEqual(status, "suspicious")
        self.assertIn("face_too_small", issues)


class DirectCoordinatorHierarchyTests(unittest.IsolatedAsyncioTestCase):
    async def test_all_participant_roles_report_directly_to_project_coordinator(self):
        with patch("app.db.database.get_coordinator_by_id") as get_coordinator:
            get_coordinator.return_value = object()
            for role in ("FS", "LS", "LSS"):
                self.assertEqual(
                    await _validated_hierarchy(role, "PC001", None),
                    ("PC001", None),
                )

    async def test_field_supervisor_parent_assignment_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "directly to a project coordinator"):
            await _validated_hierarchy("LS", "PC001", "EMP001")

    async def test_legacy_role_chains_migrate_to_direct_coordinator_assignments(self):
        participants = [
            {"employee_id": "FS001", "role": "FS", "pc_id": "PC001", "parent_employee_id": "OLD"},
            {"employee_id": "LS001", "role": "LS", "pc_id": None, "parent_employee_id": "FS001"},
            {"employee_id": "LSS001", "role": "LSS", "pc_id": None, "parent_employee_id": "LS001"},
        ]
        parent_entries = {
            "FS001": SimpleNamespace(pc_id="PC001", parent_employee_id=None),
            "LS001": SimpleNamespace(pc_id=None, parent_employee_id="FS001"),
        }

        class AsyncDocuments:
            async def __aiter__(self):
                for participant in participants:
                    yield participant

        database = SimpleNamespace(
            roster=SimpleNamespace(
                find=lambda _query: AsyncDocuments(),
                update_one=AsyncMock(),
            )
        )
        with (
            patch("app.db.database.get_db", return_value=database),
            patch(
                "app.db.database.get_roster_by_employee_id",
                new_callable=AsyncMock,
                side_effect=lambda employee_id: parent_entries.get(employee_id),
            ),
            patch(
                "app.db.database.get_coordinator_by_id",
                new_callable=AsyncMock,
                side_effect=lambda pc_id: object() if pc_id == "PC001" else None,
            ),
        ):
            migrated = await migrate_legacy_reporting_hierarchy()

        self.assertEqual(migrated, 3)
        self.assertEqual(database.roster.update_one.await_count, 3)
        for call in database.roster.update_one.await_args_list:
            self.assertEqual(call.args[1], {"$set": {"pc_id": "PC001", "parent_employee_id": None}})


class AttendanceExportTests(unittest.TestCase):
    def test_workbook_includes_coordinator_and_serial_numbered_attendance_sheets(self):
        workbook_bytes = build_export_workbook(
            coordinators=[{
                "pc_id": "PC001",
                "name": "Coordinator",
                "phone_number": "911234567890",
                "is_active": True,
            }],
            employees=[{
                "employee_id": "EMP001",
                "name": "Local Supervisor",
                "role": "LS",
                "pc_id": "PC001",
                "area": "Pune",
                "is_active": True,
            }],
            grids={"EMP001": {"2026-10-22": "P"}},
            days=["2026-10-22"],
        )
        workbook = load_workbook(BytesIO(workbook_bytes), data_only=False)

        self.assertEqual(workbook["Summary"]["B4"].value, 1)
        self.assertEqual(workbook["Summary"]["B5"].value, 1)
        self.assertEqual(workbook["Project Coordinators"]["A2"].value, 1)
        self.assertEqual(workbook["Project Coordinators"]["F2"].value, 1)
        attendance = workbook["Local supervisors"]
        self.assertEqual(attendance["A2"].value, 1)
        self.assertEqual(attendance["E2"].value, "Coordinator (PC001)")
        self.assertEqual(attendance["G2"].value, "P")
        self.assertEqual(attendance["H2"].value, '=COUNTIF(G2:G2,"P")+COUNTIF(G2:G2,"L")+COUNTIF(G2:G2,"M")')

    def test_participant_details_workbook_contains_roster_only_columns_and_total(self):
        workbook_bytes = build_participant_details_workbook(
            employees=[{
                "employee_id": "EMP002",
                "name": "Nipahd",
                "role": "LS",
                "pc_id": "PC001",
                "area": "Nipahd",
                "designation": "LS",
                "is_active": True,
            }, {
                "employee_id": "EMP001",
                "name": "Charan",
                "role": "FS",
                "pc_id": "PC001",
                "area": "Asthi",
                "designation": "Field Supervisor",
                "is_active": True,
            }],
            coordinator_names={"PC001": "Naveen"},
        )
        workbook = load_workbook(BytesIO(workbook_bytes), data_only=True)
        sheet = workbook["Participant Details"]

        self.assertEqual(
            [sheet.cell(1, column).value for column in range(1, 9)],
            ["S.No", "ID", "Name", "Department", "Reports To", "Area", "Designation", "Status"],
        )
        self.assertEqual(
            [sheet.cell(2, column).value for column in range(1, 9)],
            [1, "EMP001", "Charan", "FS", "Naveen", "Asthi", "Field Supervisor", "Active"],
        )
        self.assertEqual(sheet["A3"].value, 2)
        self.assertEqual(sheet["B3"].value, "EMP002")
        self.assertEqual(sheet["A5"].value, "Total participants")
        self.assertEqual(sheet["B5"].value, 2)
        self.assertEqual(workbook.sheetnames, ["Participant Details"])


class ParticipantDetailsExportEndpointTests(unittest.IsolatedAsyncioTestCase):
    async def test_export_applies_participant_filters_and_returns_xlsx_without_attendance(self):
        employee = RosterEntry(
            phone_number="911234567890",
            employee_id="EMP001",
            name="Charan",
            role="FS",
            pc_id="PC001",
            area="Asthi",
            designation="Field Supervisor",
        )
        coordinator = Coordinator(pc_id="PC001", name="Naveen")
        admin = AdminUser(username="admin", password_hash="hash")
        with (
            patch("app.main.list_employees", new_callable=AsyncMock, return_value=([employee], 1)) as list_people,
            patch("app.main.list_coordinators", new_callable=AsyncMock, return_value=([coordinator], 1)),
        ):
            response = await export_employee_details(
                pc_id="PC001",
                area="Asthi",
                is_active=True,
                search="Charan",
                admin=admin,
            )

        list_people.assert_awaited_once_with(
            pc_id="PC001",
            area="Asthi",
            is_active=True,
            search="Charan",
            page=1,
            page_size=100_000,
        )
        self.assertEqual(
            response.headers["content-type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(response.body), data_only=True)
        self.assertEqual(workbook.sheetnames, ["Participant Details"])
        self.assertEqual(workbook.active["E2"].value, "Naveen")
        self.assertNotIn("Summary", workbook.sheetnames)
