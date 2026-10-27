"""Create a formatted, date-wise attendance Excel workbook."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MISSING = "Not Assigned"
NOT_PROVIDED = "Not Provided"
NO_RESULT = "—"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
TITLE_FILL = PatternFill("solid", fgColor="17365D")
HEADER_FONT = Font(bold=True, color="FFFFFF")

THIN = Side(style="thin", color="D9E2F3")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MARKER_FILLS = {
    "P": PatternFill("solid", fgColor="C6EFCE"),
    "L": PatternFill("solid", fgColor="D9EAF7"),
    "AL": PatternFill("solid", fgColor="FF9999"),
    "HD": PatternFill("solid", fgColor="FFE4B5"),
    "WO": PatternFill("solid", fgColor="D9D9D9"),
    "H": PatternFill("solid", fgColor="D9EAD3"),
    "A": PatternFill("solid", fgColor="FFC7CE"),
    "F": PatternFill("solid", fgColor="FFEB9C"),
    "M": PatternFill("solid", fgColor="C6EFCE"),
    "N/A": PatternFill("solid", fgColor="D9D9D9"),
    NO_RESULT: PatternFill("solid", fgColor="F2F2F2"),
}


def _value(value, fallback=MISSING):
    if value is None:
        return fallback
    if isinstance(value, str) and not value.strip():
        return fallback
    return value


def _style_sheet(ws, header_row: int = 1):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f"A{header_row + 1}"
    for cell in ws[header_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = CELL_BORDER
    ws.row_dimensions[header_row].height = 28


def _fit_columns(ws, max_width: int = 28):
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        longest = max(
            (len(str(cell.value)) for cell in column_cells if cell.value is not None),
            default=10,
        )
        ws.column_dimensions[letter].width = min(max(longest + 2, 11), max_width)


def build_participant_details_workbook(
    employees: list[dict],
    coordinator_names: dict[str, str],
) -> bytes:
    """Create a roster-only Excel workbook without attendance information."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Participant Details"
    ws.append([
        "S.No",
        "ID",
        "Name",
        "Department",
        "Reports To",
        "Area",
        "Designation",
        "Status",
    ])

    for serial, employee in enumerate(
        sorted(employees, key=lambda item: str(item.get("name") or "").casefold()),
        start=1,
    ):
        pc_id = employee.get("pc_id")
        reports_to = coordinator_names.get(str(pc_id), pc_id) if pc_id else None
        status = "Active"
        if not employee.get("is_active", True):
            status = f"Inactive since {employee.get('deactivated_at') or '—'}"
        ws.append([
            serial,
            _value(employee.get("employee_id")),
            _value(employee.get("name")),
            _value(employee.get("role")),
            _value(reports_to),
            _value(employee.get("area")),
            _value(employee.get("designation")),
            status,
        ])

    _style_sheet(ws)
    ws.auto_filter.ref = f"A1:H{max(1, len(employees) + 1)}"
    for column, width in zip("ABCDEFGH", [8, 16, 26, 16, 28, 20, 24, 24]):
        ws.column_dimensions[column].width = width

    total_row = len(employees) + 3
    ws.cell(total_row, 1, "Total participants")
    ws.cell(total_row, 2, len(employees))
    ws.cell(total_row, 1).font = Font(bold=True)
    ws.cell(total_row, 2).font = Font(bold=True)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def build_coordinator_details_workbook(
    coordinators: list[dict],
    employees: list[dict],
) -> bytes:
    """Create a project-coordinator details workbook without attendance data."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Project Coordinators"
    ws.append(["S.No", "ID", "Name", "Phone", "Status", "Participant Areas"])

    employees_by_pc: dict[str, list[dict]] = {}
    for employee in employees:
        pc_id = employee.get("pc_id")
        if pc_id:
            employees_by_pc.setdefault(str(pc_id), []).append(employee)

    for serial, coordinator in enumerate(
        sorted(coordinators, key=lambda item: str(item.get("name") or "").casefold()),
        start=1,
    ):
        pc_id = str(coordinator.get("pc_id", ""))
        areas = sorted({
            str(employee["area"]).strip()
            for employee in employees_by_pc.get(pc_id, [])
            if employee.get("area") and str(employee["area"]).strip()
        })
        status = "Active"
        if not coordinator.get("is_active", True):
            status = f"Inactive since {coordinator.get('deactivated_at') or '—'}"
        ws.append([
            serial,
            _value(pc_id),
            _value(coordinator.get("name")),
            _value(coordinator.get("phone_number")),
            status,
            ", ".join(areas) or "—",
        ])

    _style_sheet(ws)
    ws.auto_filter.ref = f"A1:F{max(1, len(coordinators) + 1)}"
    for column, width in zip("ABCDEF", [8, 16, 26, 20, 24, 32]):
        ws.column_dimensions[column].width = width

    total_row = len(coordinators) + 3
    ws.cell(total_row, 1, "Total project co-ordinators")
    ws.cell(total_row, 2, len(coordinators))
    ws.cell(total_row, 1).font = Font(bold=True)
    ws.cell(total_row, 2).font = Font(bold=True)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def build_export_workbook(
    coordinators: list[dict],
    employees: list[dict],
    grids: dict[str, dict[str, str] | None],
    days: list[str],
) -> bytes:
    wb = Workbook()
    people = [
        {
            "employee_id": pc["pc_id"],
            "name": pc.get("name"),
            "role": "PC",
        }
        for pc in coordinators
    ] + employees

    # --- Summary Sheet ---
    summary = wb.active
    summary.title = "Summary"
    summary.sheet_view.showGridLines = False

    summary.append(["MONTHLY ATTENDANCE REPORT"])
    summary.merge_cells("A1:B1")
    summary["A1"].fill = TITLE_FILL
    summary["A1"].font = Font(bold=True, size=15, color="FFFFFF")
    summary["A1"].alignment = Alignment(vertical="center")
    summary.row_dimensions[1].height = 30

    summary.append(["Report Start", days[0] if days else NOT_PROVIDED])
    summary.append(["Report End", days[-1] if days else NOT_PROVIDED])
    summary.append([
        "Total Attendance Participants",
        len(employees),
    ])
    summary.append(["Total Project Coordinators", len(coordinators)])
    summary.append([])

    summary.append(["Code", "Meaning"])
    legend_row = summary.max_row
    for cell in summary[legend_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    legend = [
        ("P", "Present — on time"),
        ("L", "Present — late"),
        ("A", "Absent"),
        ("AL", "Absent — after grace period"),
        ("HD", "Half day"),
        ("WO", "Week off"),
        ("H", "Holiday"),
        ("F", "Flagged — awaiting review"),
        ("M", "Manual attendance entry"),
        ("N/A", "Inactive on this date"),
        (NO_RESULT, "No result yet — today or future date"),
    ]
    for marker, meaning in legend:
        summary.append([marker, meaning])
        summary.cell(summary.max_row, 1).fill = MARKER_FILLS.get(marker, PatternFill())

    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 50

    # --- Project coordinator worksheet ---
    coordinator_sheet = wb.create_sheet("Project Coordinators")
    coordinator_sheet.append([
        "S.No",
        "Coordinator ID",
        "Name",
        "Phone",
        "Status",
        "Attendance Participants",
        "Areas",
    ])
    participants_by_pc: dict[str, list[dict]] = {}
    for employee in employees:
        pc_id = employee.get("pc_id")
        if pc_id:
            participants_by_pc.setdefault(pc_id, []).append(employee)
    for serial, coordinator in enumerate(
        sorted(coordinators, key=lambda item: str(item.get("name") or "").casefold()),
        start=1,
    ):
        pc_id = str(coordinator.get("pc_id", ""))
        people_for_pc = participants_by_pc.get(pc_id, [])
        areas = sorted({
            str(person["area"]).strip()
            for person in people_for_pc
            if person.get("area") and str(person["area"]).strip()
        })
        coordinator_sheet.append([
            serial,
            pc_id,
            _value(coordinator.get("name")),
            _value(coordinator.get("phone_number")),
            "Active" if coordinator.get("is_active", True) else "Inactive",
            len(people_for_pc),
            ", ".join(areas) or "—",
        ])
    _style_sheet(coordinator_sheet)
    coordinator_sheet.freeze_panes = "A2"
    coordinator_sheet.auto_filter.ref = coordinator_sheet.dimensions
    coordinator_sheet.column_dimensions["A"].width = 8
    coordinator_sheet.column_dimensions["B"].width = 18
    coordinator_sheet.column_dimensions["C"].width = 28
    coordinator_sheet.column_dimensions["D"].width = 18
    coordinator_sheet.column_dimensions["E"].width = 14
    coordinator_sheet.column_dimensions["F"].width = 26
    coordinator_sheet.column_dimensions["G"].width = 40

    # --- Department attendance worksheets ---
    roles = [
        ("FS", "Field Supervisors"),
        ("LS", "Local supervisors"),
        ("LSS", "Local supervisor scanning"),
    ]
    person_by_id = {str(person["employee_id"]): person for person in people}
    for role, sheet_title in roles:
        ws_pc = wb.create_sheet(sheet_title)
        # Editable daily codes; totals recalculate automatically in Excel.
        headers = ["S.No", "ID", "Location", "Name", "Reports To", "Position"] + days + ["Present", "Absent", "HD", "WO", "Holidays", "Late", "AL"]
        ws_pc.append(headers)
        from openpyxl.worksheet.datavalidation import DataValidation
        validation = DataValidation(type="list", formula1='"P,L,A,AL,HD,WO,H,M,F,N/A"', allow_blank=True)
        validation.error = "Select an attendance code from the list"
        validation.showErrorMessage = True
        ws_pc.add_data_validation(validation)
        role_people = [
            person for person in people
            if person.get("role", "FS") == role
            and grids.get(str(person.get("employee_id", ""))) is not None
        ]
        for serial, emp in enumerate(
            sorted(role_people, key=lambda x: str(x.get("name") or "").casefold()),
            start=1,
        ):
            emp_id = str(emp.get("employee_id", ""))
            grid = grids.get(emp_id) or {}
            markers = [_value(grid.get(day), NO_RESULT) for day in days]
            site = emp.get("assigned_site") or {}
            location = site.get("name") if isinstance(site, dict) else None
            parent_id = emp.get("parent_employee_id") or emp.get("pc_id")
            parent = person_by_id.get(str(parent_id)) if parent_id else None
            parent_label = f"{parent.get('name')} ({parent_id})" if parent else _value(parent_id, "—")
            position = emp.get("designation") or {
                "FS": "Field Supervisor",
                "LS": "Local Supervisor",
                "LSS": "Local Supervisor Scanning",
            }[role]
            ws_pc.append([
                serial,
                _value(emp.get("employee_id")),
                _value(location or emp.get("area")),
                _value(emp.get("name")),
                parent_label,
                position,
            ] + markers)
            r = ws_pc.max_row
            first = 7
            last = 6 + len(days)
            range_ref = f"{get_column_letter(first)}{r}:{get_column_letter(last)}{r}"
            formulas = [
                f'=COUNTIF({range_ref},"P")+COUNTIF({range_ref},"L")+COUNTIF({range_ref},"M")',
                f'=COUNTIF({range_ref},"A")+COUNTIF({range_ref},"AL")',
                f'=COUNTIF({range_ref},"HD")', f'=COUNTIF({range_ref},"WO")',
                f'=COUNTIF({range_ref},"H")', f'=COUNTIF({range_ref},"L")', f'=COUNTIF({range_ref},"AL")',
            ]
            for offset, formula in enumerate(formulas, start=last + 1):
                ws_pc.cell(r, offset, formula)
            for c, marker in enumerate(markers, start=7):
                cell = ws_pc.cell(r, c)
                cell.fill = MARKER_FILLS.get(marker, PatternFill())
                cell.alignment = Alignment(horizontal="center")
                validation.add(cell)
        _style_sheet(ws_pc)
        ws_pc.freeze_panes = "G2"
        ws_pc.auto_filter.ref = ws_pc.dimensions
        ws_pc.column_dimensions["A"].width = 8
        ws_pc.column_dimensions["B"].width = 14
        ws_pc.column_dimensions["C"].width = 24
        ws_pc.column_dimensions["D"].width = 26
        ws_pc.column_dimensions["E"].width = 32
        ws_pc.column_dimensions["F"].width = 28
        for col_idx in range(7, 7 + len(days)):
            ws_pc.column_dimensions[get_column_letter(col_idx)].width = 13
        for col_idx in range(7 + len(days), 14 + len(days)):
            ws_pc.column_dimensions[get_column_letter(col_idx)].width = 14

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
